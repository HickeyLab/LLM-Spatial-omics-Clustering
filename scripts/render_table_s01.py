#!/usr/bin/env python3
"""Build an editable, presentation-only Table S1 from current manuscript cells."""
from __future__ import annotations

import argparse
import hashlib
from html import escape
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def load_inputs():
    config = json.loads((ROOT / "configs/table_s01.json").read_text())
    path = ROOT / config["data"]
    if hashlib.sha256(path.read_bytes()).hexdigest() != config["data_sha256"]:
        raise ValueError("Table S1 frozen text checksum differs")
    data = json.loads(path.read_text())
    if len(data["headers"]) != 5 or len(data["rows"]) != 8:
        raise ValueError("Table S1 must have five columns and eight body rows")
    if any(len(row) != 5 for row in data["rows"]):
        raise ValueError("Unexpected Table S1 row width")
    if abs(sum(config["column_width_inches"]) - 6.5) > 1e-8:
        raise ValueError("Table S1 must fit the manuscript's 6.5-inch text width")
    return config, data


def element(parent, tag, **attributes):
    node = OxmlElement(f"w:{tag}")
    for key, value in attributes.items():
        node.set(qn(f"w:{key}"), str(value))
    parent.append(node)
    return node


def build_document(config, data):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    normal = document.styles["Normal"]
    normal.font.name = config["font"]
    normal.font.size = Pt(config["font_size_pt"])
    normal.font.color.rgb = RGBColor.from_string(config["text_color"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = config["line_spacing"]
    table = document.add_table(rows=9, cols=5)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(table.columns, config["column_width_inches"]):
        col.width = Inches(width)
    properties = table._tbl.tblPr
    borders = element(properties, "tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        element(borders, side, val="single", sz=4, color=config["border_color"])
    matrix = [data["headers"], *data["rows"]]
    for row_index, (row, values) in enumerate(zip(table.rows, matrix)):
        row_properties = row._tr.get_or_add_trPr()
        element(row_properties, "cantSplit")
        if row_index == 0:
            element(row_properties, "tblHeader")
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, config["column_width_inches"])):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            padding = element(cell_properties, "tcMar")
            for side, amount in config["cell_padding_twips"].items():
                element(padding, side, w=amount, type="dxa")
            fill = config["header_fill"] if row_index == 0 else config["field_fill"] if col_index == 0 else config["body_fill"]
            element(cell_properties, "shd", val="clear", fill=fill)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.keep_together = True
            short_value = row_index in (1, 4, 8) and col_index > 0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_index == 0 and col_index > 0) or short_value else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            run.bold = row_index == 0 or col_index == 0
            run.font.name = config["font"]
            run.font.size = Pt(config["header_font_size_pt"] if row_index == 0 else config["font_size_pt"])
            run.font.color.rgb = RGBColor.from_string("000000" if row_index == 0 else config["text_color"])
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.keep_together = True
    lead, separator, rest = data["caption"].partition("benchmark.")
    if not separator:
        raise ValueError("Expected Table S1 caption lead")
    run = paragraph.add_run(lead + separator)
    run.bold = True
    paragraph.add_run(rest)
    for run in paragraph.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string("000000")
    document.core_properties.title = "Table S1 clustering configurations"
    document.core_properties.subject = "Presentation revision preserving current manuscript table content"
    return document


def html_table(config, data):
    """Retain a native editable-table path for pasting into the existing document."""
    rows = []
    for r, values in enumerate([data["headers"], *data["rows"]]):
        cells = []
        for c, (value, width) in enumerate(zip(values, config["column_width_inches"])):
            fill = config["header_fill"] if r == 0 else config["field_fill"] if c == 0 else config["body_fill"]
            weight = "bold" if r == 0 or c == 0 else "normal"
            align = "center" if (r == 0 and c > 0) or (r in (1, 4, 8) and c > 0) else "left"
            style = (f"width:{width}in;padding:5pt 4.5pt;border:0.5pt solid #{config['border_color']};"
                     f"background-color:#{fill};font-family:{config['font']};font-size:{config['header_font_size_pt'] if r == 0 else config['font_size_pt']}pt;"
                     f"font-weight:{weight};color:#{config['text_color']};text-align:{align};vertical-align:middle;line-height:108%;")
            tag = "th" if r == 0 else "td"
            cells.append(f'<{tag} style="{style}">{escape(value)}</{tag}>')
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return '<table style="border-collapse:collapse;table-layout:fixed;width:6.5in;">' + "".join(rows) + "</table>\n"


def verify_document(path, data):
    document = Document(path)
    if len(document.tables) != 1:
        raise ValueError("Expected exactly one Table S1")
    actual = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
    if actual != [data["headers"], *data["rows"]]:
        raise ValueError("A Table S1 cell changed")
    if document.paragraphs[-1].text != data["caption"]:
        raise ValueError("Table S1 caption changed")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs/source_rebuilt/table_s01")
    parser.add_argument("--check", type=Path, help="Verify an existing editable DOCX's content")
    args = parser.parse_args()
    config, data = load_inputs()
    if args.check:
        verify_document(args.check, data)
        print("Verified all 45 Table S1 cells and the unchanged caption")
        return
    args.output_dir.mkdir(parents=True, exist_ok=True)
    target = args.output_dir / "Table_S1_professional.docx"
    build_document(config, data).save(target)
    (args.output_dir / "Table_S1_professional.html").write_text(html_table(config, data))
    verify_document(target, data)
    print(target)


if __name__ == "__main__":
    main()
